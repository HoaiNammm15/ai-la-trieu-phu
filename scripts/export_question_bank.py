"""Export the game's canonical question bank to Excel and UTF-8 CSV.

Requires openpyxl. Run: python scripts/export_question_bank.py
"""
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['id', 'level', 'question', 'a', 'b', 'c', 'd', 'correct', 'category', 'explanation', 'active']
STAGES = [(1, 5, 'Dễ', 'E4F3E8'), (6, 10, 'Trung bình', 'FFF0CC'), (11, 15, 'Khó', 'FCE2E2')]
XLSX_PATH = ROOT / 'docs' / '1000-cau-hoi-ai-la-trieu-phu.xlsx'
CSV_PATH = ROOT / 'docs' / 'questions-1000.csv'


def rows_for(bank):
    return [[q['id'], q['level'], q['question'], *q['answers'], 'ABCD'[q['correct']],
             q['category'], q['explanation'], True] for q in bank]


def style_sheet(sheet, widths, height=46):
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17324D')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = height
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = 's'
                cell.number_format = '@'
    sheet.auto_filter.ref = sheet.dimensions


def export():
    bank = json.loads((ROOT / 'src/data/questions-bank.json').read_text(encoding='utf-8'))
    review = json.loads((ROOT / 'docs/questions-review.json').read_text(encoding='utf-8'))
    assert len(bank) == 1000, 'The canonical bank must contain exactly 1000 questions.'
    data = rows_for(bank)
    with CSV_PATH.open('w', encoding='utf-8-sig', newline='') as output:
        writer = csv.writer(output)
        writer.writerow(HEADERS)
        for row in data:
            writer.writerow([*row[:-1], 'TRUE'])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Trang tính1'
    sheet.append(HEADERS)
    for row in data:
        sheet.append(row)
    style_sheet(sheet, [9, 9, 72, 28, 28, 28, 28, 12, 22, 85, 12], height=62)
    sheet.freeze_panes = 'C2'
    sheet.sheet_properties.tabColor = '17324D'
    for row in sheet.iter_rows(min_row=2):
        color = next(color for low, high, _, color in STAGES if low <= row[1].value <= high)
        row[1].fill = PatternFill('solid', fgColor=color)
        row[1].font = Font(name='Calibri', size=11, bold=True)
        row[7].font = Font(name='Calibri', size=11, bold=True, color='146B3A')
        for index in [0, 1, 7, 10]:
            row[index].alignment = Alignment(horizontal='center', vertical='top')
    table = Table(displayName='QuestionBank', ref='A1:K1001')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    for column, validation in [
        ('B', DataValidation(type='whole', operator='between', formula1=1, formula2=15)),
        ('H', DataValidation(type='list', formula1='"A,B,C,D"')),
        ('K', DataValidation(type='list', formula1='"TRUE,FALSE"')),
    ]:
        validation.allow_blank = False
        validation.errorTitle = 'Giá trị không hợp lệ'
        validation.error = 'Dùng level 1–15, đáp án A–D và active TRUE/FALSE theo cột.'
        validation.showErrorMessage = True
        validation.errorStyle = 'stop'
        sheet.add_data_validation(validation)
        validation.add(f'{column}2:{column}5000')
    sheet.print_title_rows = '1:1'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    guide = workbook.create_sheet('Hướng dẫn')
    guide.append(['Nội dung', 'Quy định / kết quả'])
    guide_rows = [
        ['Bộ câu hỏi', '1.000 câu đang bật: giữ ID 1–170 từ file gốc, bổ sung ID 171–1000.'],
        ['Mốc 1: câu 1–5', 'Lấy ngẫu nhiên 5 câu thuộc level 1–5 (dễ), hướng tới mốc 2 triệu.'],
        ['Mốc 2: câu 6–10', 'Lấy ngẫu nhiên 5 câu thuộc level 6–10 (trung bình), hướng tới mốc 22 triệu.'],
        ['Mốc 3: câu 11–15', 'Lấy ngẫu nhiên 5 câu thuộc level 11–15 (khó), hướng tới đích 150 triệu.'],
        ['Thứ tự trong mỗi mốc', 'Các câu đã bốc được xếp theo level tăng dần; không bắt buộc mỗi level đúng một câu.'],
        ['Ý nghĩa của level', 'Độ khó biên tập theo kiến thức phổ thông; chưa hiệu chỉnh bằng thống kê người chơi.'],
        ['Điều kiện khởi tạo lượt', 'Mọi câu đang bật phải có level 1–15; mỗi nhóm cần ít nhất 5 câu.'],
        ['Đáp án', 'Bốn lựa chọn riêng biệt; correct là A/B/C/D và explanation giải thích đáp án.'],
        ['ID', 'ID duy nhất. Giữ ID khi sửa nội dung; level không phải ID hoặc số thứ tự trong lượt.'],
        ['Câu dùng địa giới cũ', 'ID 33 và 80 ghi rõ mốc năm 2024, trước đợt sắp xếp hành chính năm 2025.'],
        ['Nhập Google Sheets', "Nhập tab Trang tính1, giữ 11 cột; range 'Trang tính1'!A1:K5000."],
        ['Nguồn game trên máy', 'Khi chưa đặt GOOGLE_SHEET_ID, API dùng đủ 1.000 câu trong questions-bank.json.'],
        ['Nguồn Google Sheets', 'Khi có GOOGLE_SHEET_ID, API dùng dữ liệu ở Sheet đó; cần nhập bản Excel/CSV này để cập nhật.'],
        ['Kiểm tra nội dung', 'Đã biên tập câu hỏi gốc và câu bổ sung; các nguồn tra cứu cụ thể nằm ở tab Nguồn tham khảo.'],
    ]
    for row in guide_rows:
        guide.append(row)
    style_sheet(guide, [35, 120])

    stats = workbook.create_sheet('Thống kê')
    stats.append(['Phân loại', 'Nhóm', 'Số câu'])
    for low, high, name, _ in STAGES:
        stats.append(['Mốc chơi', f'{name}: level {low}–{high}', sum(low <= q['level'] <= high for q in bank)])
    for level, count in sorted(Counter(q['level'] for q in bank).items()):
        stats.append(['Level', level, count])
    for category, count in sorted(Counter(q['category'] for q in bank).items()):
        stats.append(['Chủ đề', category, count])
    for answer, count in sorted(Counter('ABCD'[q['correct']] for q in bank).items()):
        stats.append(['Đáp án đúng', answer, count])
    stats.append(['Tổng', 'Câu đang bật', 1000])
    style_sheet(stats, [24, 45, 15], height=25)

    changes = workbook.create_sheet('Rà soát câu gốc')
    changes.append(['ID', 'Level cũ', 'Level mới', 'Nội dung chỉnh', 'Câu hỏi sau rà soát'])
    for item in review['original_review']:
        changes.append([item['id'], item['old_level'], item['level'], item['reason'], bank[item['id'] - 1]['question']])
    style_sheet(changes, [9, 12, 12, 90, 85], height=62)

    sources = workbook.create_sheet('Nguồn tham khảo')
    sources.append(['Phạm vi', 'Nguồn / ghi chú', 'URL'])
    for item in review.get('sources', []):
        sources.append([str(item.get('ids', item.get('scope', 'Kiến thức tham khảo'))),
                        item.get('title', item.get('note', item.get('claim', 'Nguồn tra cứu'))), item['url']])
    style_sheet(sources, [30, 90, 100], height=50)
    for row in sources.iter_rows(min_row=2):
        row[2].hyperlink = row[2].value
        row[2].font = Font(name='Calibri', size=11, color='0563C1', underline='single')

    workbook.save(XLSX_PATH)
    # Verify the artifacts match the canonical game data, including date strings.
    saved = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    assert [list(r) for r in saved['Trang tính1'].values] == [HEADERS, *data]
    saved.close()
    with CSV_PATH.open(encoding='utf-8-sig', newline='') as exported:
        parsed = list(csv.reader(exported))
    assert parsed[0] == HEADERS and len(parsed) == 1001
    assert parsed[1:] == [[str(v) for v in row[:-1]] + ['TRUE'] for row in data]
    print(f'Excel: {XLSX_PATH}\nCSV: {CSV_PATH}\nVerified 1,000 rows against game JSON.')


if __name__ == '__main__':
    export()
